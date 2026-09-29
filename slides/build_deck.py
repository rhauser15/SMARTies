"""Build the Marriott x Agentforce demo deck.

    python3 -m pip install python-pptx
    python3 slides/build_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

RED = RGBColor(0xB4, 0x1F, 0x3A)
RED_DARK = RGBColor(0x7A, 0x14, 0x28)
INK = RGBColor(0x1C, 0x1C, 0x1C)
CHAR = RGBColor(0x26, 0x20, 0x22)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
GOLD = RGBColor(0xB3, 0x8E, 0x5D)
BLUSH = RGBColor(0xFB, 0xF1, 0xF3)
LINE = RGBColor(0xE6, 0xE1, 0xDC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xD9, 0xCF, 0xD1)

HEAD = "Cambria"
BODY = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        segs = p if isinstance(p, list) else [(p, {})]
        for seg, o in segs:
            r = para.add_run()
            r.text = seg
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", italic)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill, radius=True, line=None, shadow=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if not shadow:
        shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def badge(slide, x, y, d, label, fill=RED, color=WHITE, size=20):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = HEAD, Pt(size), True, color
    return c


def wordmark(slide, x, y, dark=False):
    col = WHITE if dark else INK
    text(slide, x, y, 3, 0.45, "Marriott", size=22, font=HEAD, bold=True, color=col)
    text(slide, x + 0.02, y + 0.42, 3, 0.25, "BONVOY  ×  SALESFORCE", size=9, bold=True,
         color=SOFT if dark else MUTED)


def section_tag(slide, label, x=0.6, y=0.45, dark=False):
    text(slide, x, y, 6, 0.3, label.upper(), size=11, bold=True, color=GOLD if not dark else GOLD)


def footer(slide, n, dark=False):
    text(slide, 0.6, 7.0, 8, 0.25, "Mock scenario for Salesforce SE onboarding · Not affiliated with Marriott International",
         size=9, color=SOFT if dark else MUTED)
    text(slide, 12.2, 7.0, 0.55, 0.25, str(n), size=9, color=SOFT if dark else MUTED, align=PP_ALIGN.RIGHT)


def timing(slide, label, dark=False):
    b = box(slide, 11.35, 0.4, 1.4, 0.42, RED if not dark else WHITE)
    tf = b.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = BODY, Pt(12), True
    r.font.color.rgb = WHITE if not dark else RED


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
bg(s, CHAR)
panel = box(s, 7.9, 0, 5.433, 7.5, RED, radius=False)
wordmark(s, 0.8, 0.7, dark=True)
text(s, 0.8, 2.2, 6.8, 0.4, "AGENTFORCE FOR THE CUSTOMER ENGAGEMENT CENTER", size=12, bold=True, color=GOLD)
text(s, 0.8, 2.65, 6.9, 2.2, "Where are my points?", size=54, font=HEAD, bold=True, color=WHITE)
text(s, 0.8, 4.75, 6.6, 1.2,
     "How Marriott can resolve its #1 guest request instantly, and give every associate the whole story.",
     size=18, color=SOFT)
text(s, 0.8, 6.3, 6.8, 0.4, "[Team name]  ·  Solutions Foundations  ·  2026", size=13, color=SOFT)
# right panel: three big numbers as a visual
for i, (num, lbl) in enumerate([("30M", "cases a year"), ("35%", "are missing stays & points"), ("7", "global contact centers")]):
    y = 1.35 + i * 1.75
    text(s, 8.6, y, 4.2, 0.9, num, size=54, font=HEAD, bold=True, color=WHITE)
    text(s, 8.6, y + 0.95, 4.2, 0.4, lbl, size=15, color=WHITE)
s.notes_slide.notes_text_frame.text = (
    "OPENING (0:00). Walk up, pause, smile. Title stays up while the team introduces itself.")

# ---------------------------------------------------------------- 2. Credential
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
section_tag(s, "Opening · Credential")
timing(s, "0:00 – 0:30")
text(s, 0.6, 0.8, 11, 0.9, "Who we are and why we're here", size=38, font=HEAD, bold=True)
people = [("A", "[Name]", "Solution Engineer · Service", "Opens, credentials the team"),
          ("B", "[Name]", "Solution Engineer · Data 360", "Sets the scene"),
          ("C", "[Name]", "Solution Engineer · Agentforce", "Drives the live demo"),
          ("D", "[Name]", "Solution Engineer · Platform", "Closes, value and next steps")]
for i, (ini, name, role, part) in enumerate(people):
    x = 0.6 + i * 3.08
    box(s, x, 2.05, 2.85, 2.9, BLUSH)
    badge(s, x + 0.3, 2.35, 0.8, ini, size=22)
    text(s, x + 0.3, 3.35, 2.4, 0.45, name, size=20, font=HEAD, bold=True)
    text(s, x + 0.3, 3.8, 2.4, 0.5, role, size=13, color=MUTED)
    text(s, x + 0.3, 4.3, 2.4, 0.5, part, size=13, color=RED, bold=True)
box(s, 0.6, 5.35, 12.13, 1.3, CHAR)
text(s, 0.95, 5.5, 11.5, 1.0,
     [[("Why we're here: ", {"bold": True, "color": GOLD}),
       ("you told us your guests expect to be known, your associates are juggling too many systems, and cost to serve keeps rising. "
        "In the next 10 minutes we'll show you what that looks like fixed, live.", {"color": WHITE})]],
     size=16, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)
s.notes_slide.notes_text_frame.text = (
    "CREDENTIAL (30 sec). Each person: 1 sentence. Job without jargon, background, why here.\n"
    "Example: 'My role is to understand Marriott's priorities and match them to what Salesforce can do. "
    "I've spent X years helping service teams…'\nHandoff: 'Let me start with a number.'")

# ---------------------------------------------------------------- 3. Attention grabber
s = prs.slides.add_slide(BLANK)
bg(s, RED)
section_tag(s, "Opening · Attention grabber", dark=True)
timing(s, "0:30 – 1:15", dark=True)
text(s, 0.6, 1.3, 12, 2.4, "10.5 million", size=120, font=HEAD, bold=True, color=WHITE)
text(s, 0.6, 3.55, 11.5, 1.0,
     "times a year, a Marriott guest asks the same question:", size=28, color=WHITE)
text(s, 0.6, 4.35, 11.5, 1.0, "“Where are my points?”", size=44, font=HEAD, italic=True, bold=True, color=WHITE)
box(s, 0.6, 5.75, 8.2, 0.85, RED_DARK)
text(s, 0.9, 5.75, 7.8, 0.85,
     "35% of ~30M CEC cases are Missing Night Stays & Points, filed by web form and worked by hand.",
     size=15, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 3, dark=True)
s.notes_slide.notes_text_frame.text = (
    "ATTENTION GRABBER (45 sec). Say the number, then PAUSE for 2 seconds.\n"
    "'Ten and a half million times a year, one of your most loyal guests asks: where are my points? "
    "Today that means a web form, a 5 to 7 day wait, and an associate searching three systems for the answer.'\n"
    "Optional: add a 'State of Service' stat here (salesforce.com → State of Service report).\n"
    "Handoff: 'And that one question tells the story of three bigger challenges.'")

# ---------------------------------------------------------------- 4. Set the scene: top 3 challenges
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
section_tag(s, "Opening · Setting the scene")
timing(s, "1:15 – 2:15")
text(s, 0.6, 0.8, 11, 0.9, "Three challenges we heard", size=38, font=HEAD, bold=True)
cards = [
    ("1", "Associates swivel-chair", "30+ brands, many systems, and no single view of the guest. Tom spends his time searching instead of serving.",
     "Heard from: Tom & Mateo"),
    ("2", "Cost to serve keeps climbing", "Phone, web form and email are the only channels. Repetitive cases like missing stays drive up handle time and headcount.",
     "Heard from: Suki"),
    ("3", "Loyalty is on the line", "CSAT is trending down while Airbnb and Vrbo compete for the same traveler. Slow answers erode Bonvoy loyalty.",
     "Heard from: Sofia"),
]
for i, (n, title, body, who) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 2.0, 3.85, 4.5, BLUSH)
    badge(s, x + 0.35, 2.35, 0.85, n, size=26)
    text(s, x + 0.35, 3.45, 3.2, 0.9, title, size=22, font=HEAD, bold=True)
    text(s, x + 0.35, 4.35, 3.2, 1.6, body, size=15, color=INK)
    text(s, x + 0.35, 5.95, 3.2, 0.35, who, size=12, bold=True, color=RED)
footer(s, 4)
s.notes_slide.notes_text_frame.text = (
    "SET THE SCENE (1 min). Use their words. Name the stakeholder behind each challenge.\n"
    "Ask one open question: 'Suki, is that 35% number still what you're seeing?'\n"
    "Handoff: 'And if nothing changes…'")

# ---------------------------------------------------------------- 5. Cost of doing nothing
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
section_tag(s, "Opening · Setting the scene")
timing(s, "2:15 – 3:00")
text(s, 0.6, 0.8, 11, 0.9, "The cost of doing nothing", size=38, font=HEAD, bold=True)
# left: "today" process
text(s, 0.6, 1.95, 5.8, 0.4, "A missing-stay request today", size=16, bold=True, color=RED)
steps = ["Guest fills out a web form", "Case lands in a queue for days", "Associate searches PMS, loyalty and email",
         "Points posted by hand, with an email reply", "Guest waits 5–7 business days"]
for i, st in enumerate(steps):
    y = 2.5 + i * 0.8
    badge(s, 0.6, y, 0.5, str(i + 1), fill=CHAR, size=14)
    text(s, 1.3, y, 5.0, 0.5, st, size=16, anchor=MSO_ANCHOR.MIDDLE)
# right: stat callouts
stats = [("~10.5M", "manual missing-stay cases every year"),
         ("5–7 days", "for a guest to see their points"),
         ("CSAT", "trending down across several brands")]
for i, (num, lbl) in enumerate(stats):
    y = 1.95 + i * 1.6
    box(s, 7.1, y, 5.63, 1.35, CHAR if i == 0 else BLUSH)
    text(s, 7.45, y + 0.12, 2.6, 1.1, num, size=34, font=HEAD, bold=True,
         color=WHITE if i == 0 else RED, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.0, y + 0.12, 2.55, 1.1, lbl, size=14, color=WHITE if i == 0 else INK, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 5)
s.notes_slide.notes_text_frame.text = (
    "COST OF INACTION (45 sec). Walk the 5 steps quickly. Stress: every step is a system hop.\n"
    "'Every one of these cases costs associate time, and every day of waiting costs loyalty.'\n"
    "Handoff to the demo: 'Let us show you a different version of Lauren's week.'")

# ---------------------------------------------------------------- 6. Demo agenda
s = prs.slides.add_slide(BLANK)
bg(s, CHAR)
section_tag(s, "Story-based demo · Live", dark=True)
timing(s, "3:00 – 9:00", dark=True)
text(s, 0.6, 0.8, 11, 0.9, "One guest. Three chapters.", size=38, font=HEAD, bold=True, color=WHITE)
chapters = [
    ("1", "The guest helps herself", "Lauren, Platinum Elite",
     "Bonvoy Assistant (Agentforce) finds the Austin stay and credits 3 nights and 8,450 points in under a minute.",
     "Resolve instantly, 24/7"),
    ("2", "The associate has the whole story", "Tom, CEC Associate",
     "A Kyoto stay needs the hotel. Agentforce hands off with a summary and Guest 360, and drafts Tom's reply.",
     "One view of the guest"),
    ("3", "Leaders get control and speed", "Suki, Mateo & Jess",
     "Plain-English agent builder, trusted guardrails, and live metrics on resolution, handle time and CSAT.",
     "Built for scale, fast"),
]
for i, (n, title, who, body, msg) in enumerate(chapters):
    x = 0.6 + i * 4.1
    box(s, x, 2.0, 3.85, 4.55, RGBColor(0x36, 0x2E, 0x30))
    badge(s, x + 0.35, 2.3, 0.8, n, size=24)
    text(s, x + 1.35, 2.3, 2.3, 0.8, "~2 min", size=13, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.35, 3.3, 3.2, 0.9, title, size=21, font=HEAD, bold=True, color=WHITE)
    text(s, x + 0.35, 4.15, 3.2, 0.35, who, size=13, bold=True, color=GOLD)
    text(s, x + 0.35, 4.55, 3.2, 1.3, body, size=14, color=SOFT)
    text(s, x + 0.35, 5.95, 3.2, 0.4, "→ " + msg, size=14, bold=True, color=WHITE)
footer(s, 6, dark=True)
s.notes_slide.notes_text_frame.text = (
    "DEMO AGENDA (15 sec max), then switch to the browser.\n"
    "Tab 1: mock Bonvoy site (Lauren) · Tab 2: Service Console (Tom) · Tab 3: Agentforce Builder · Tab 4: dashboard (Suki).\n"
    "Full click path and talk track are in README.md §2.\n"
    "Each chapter: CONTEXT → FEATURE (live) → BENEFIT.")

# ---------------------------------------------------------------- 7. Value summary
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
section_tag(s, "Close · Value summary")
timing(s, "9:00 – 9:25")
text(s, 0.6, 0.8, 11, 0.9, "What this means for Marriott", size=38, font=HEAD, bold=True)
msgs = [("Resolve instantly, 24/7", "The #1 request is handled end to end with no form, no queue and no wait."),
        ("One view of the guest", "Every handoff carries the full story, so associates serve instead of searching."),
        ("Built for scale, fast", "Low-code, trusted, and connected to the systems Marriott already runs.")]
for i, (h, b) in enumerate(msgs):
    y = 2.0 + i * 1.5
    badge(s, 0.6, y, 0.7, str(i + 1), size=20)
    text(s, 1.55, y - 0.02, 4.9, 0.45, h, size=20, font=HEAD, bold=True)
    text(s, 1.55, y + 0.45, 4.9, 0.8, b, size=14, color=MUTED)
cd = CategoryChartData()
cd.categories = ["Today", "With Agentforce"]
cd.add_series("Missing-stay cases worked by associates (M/yr)", (10.5, 4.2))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.9), Inches(1.85), Inches(5.8), Inches(4.3), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Missing-stay cases worked by associates (M / yr)"
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size, tp.runs[0].font.name, tp.runs[0].font.color.rgb = Pt(13), BODY, INK
plot = ch.plots[0]
plot.gap_width = 70
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format, dl.number_format_is_linked = '0.0"M"', False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(16), True, INK
ser = plot.series[0]
for idx, col in enumerate([CHAR, RED]):
    pt = ser.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
va = ch.value_axis
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = 12
ca = ch.category_axis
ca.tick_labels.font.size, ca.tick_labels.font.color.rgb = Pt(13), INK
ca.format.line.color.rgb = LINE
text(s, 6.9, 6.2, 5.8, 0.6,
     "Illustrative: 60% agent resolution → ~6.3M cases/yr that never reach an associate. "
     "At a $5–6 assisted-contact cost, that is ~$30–38M/yr in capacity. Validate with Marriott.",
     size=11, color=MUTED, italic=True)
footer(s, 7)
s.notes_slide.notes_text_frame.text = (
    "VALUE SUMMARY (25 sec). Repeat the three key messages word for word from the demo.\n"
    "Frame the number as a hypothesis: 'If Agentforce resolved 60% of missing-stay requests, that's six million cases "
    "your associates never have to touch. We'd love to build that model with your real numbers.'")

# ---------------------------------------------------------------- 8. Customer story
s = prs.slides.add_slide(BLANK)
bg(s, BLUSH)
section_tag(s, "Close · Evidence")
timing(s, "9:25 – 9:45")
text(s, 0.6, 0.8, 11, 0.9, "Travel brands are already doing this", size=38, font=HEAD, bold=True)
box(s, 0.6, 2.0, 7.3, 4.6, WHITE)
text(s, 1.0, 2.3, 6.6, 0.4, "[CUSTOMER LOGO / NAME]", size=14, bold=True, color=RED)
text(s, 1.0, 2.85, 6.6, 2.0,
     "“[Approved customer quote about Agentforce resolving routine travel requests and freeing associates for "
     "high-touch moments.]”", size=22, font=HEAD, italic=True)
text(s, 1.0, 5.05, 6.6, 0.4, "[Name, Title, Company]", size=14, color=MUTED)
text(s, 1.0, 5.6, 6.6, 0.7,
     "Source: Slackbot “Customer Story & Deal Win Finder”. Pick a travel & hospitality Agentforce story. "
     "Replace all [brackets] before presenting.", size=11, italic=True, color=MUTED)
for i, lbl in enumerate(["[Metric 1]\ne.g. handle time", "[Metric 2]\ne.g. cases resolved by agent", "[Metric 3]\ne.g. time to launch"]):
    y = 2.0 + i * 1.6
    box(s, 8.25, y, 4.48, 1.4, RED if i == 0 else WHITE)
    num, cap = lbl.split("\n")
    text(s, 8.6, y + 0.15, 3.9, 0.65, num, size=26, font=HEAD, bold=True, color=WHITE if i == 0 else RED)
    text(s, 8.6, y + 0.8, 3.9, 0.45, cap, size=13, color=WHITE if i == 0 else MUTED)
footer(s, 8)
s.notes_slide.notes_text_frame.text = (
    "CUSTOMER STORY (20 sec). One sentence of context, one metric, one quote. Then move on.\n"
    "Find it: Slackbot → 'Customer Story & Deal Win Finder' → 'Agentforce Service travel hospitality'.")

# ---------------------------------------------------------------- 9. Next steps
s = prs.slides.add_slide(BLANK)
bg(s, CHAR)
section_tag(s, "Close · Next steps", dark=True)
timing(s, "9:45 – 10:00", dark=True)
text(s, 0.6, 0.8, 11, 0.9, "Let's prove it on one request", size=38, font=HEAD, bold=True, color=WHITE)
nexts = [("This week", "60-minute discovery workshop with Suki & Mateo", "Map the top 5 CEC requests and pull real case volumes."),
         ("Weeks 1–2", "Agentforce pilot: Missing Stays & Points", "One brand, one channel, connected to a sample of PMS stay data."),
         ("Week 3", "Business case review with Sofia & Jess", "Measured resolution rate, handle time and CSAT, plus the roll-out plan.")]
for i, (when, what, detail) in enumerate(nexts):
    x = 0.6 + i * 4.1
    box(s, x, 2.0, 3.85, 3.3, RGBColor(0x36, 0x2E, 0x30))
    text(s, x + 0.35, 2.3, 3.2, 0.4, when.upper(), size=12, bold=True, color=GOLD)
    text(s, x + 0.35, 2.8, 3.2, 1.2, what, size=20, font=HEAD, bold=True, color=WHITE)
    text(s, x + 0.35, 4.05, 3.2, 1.1, detail, size=14, color=SOFT)
box(s, 0.6, 5.65, 12.13, 1.0, RED)
text(s, 0.95, 5.65, 11.5, 1.0,
     "Can we get 60 minutes with Suki and Mateo on the calendar this week?",
     size=20, font=HEAD, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 9, dark=True)
s.notes_slide.notes_text_frame.text = (
    "NEXT STEPS (15 sec). Ask the question on the red bar directly, then stop talking and wait for the answer.\n"
    "Then: 'Thank you. We'd love your questions.'")

out = Path(__file__).with_name("Marriott_Agentforce_Demo.pptx")
prs.save(out)
print("wrote", out)
